#!/usr/bin/env python3
"""Word 专业排版：页面呼吸感 / 字体金字塔 / 行距段距 / 表格降噪 / 重点高亮。"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---- 字体金字塔（中文 / 英文·数字 强制 Arial）----
_F_T0 = "黑体"
_F_T1 = "黑体"
_F_T2 = "微软雅黑"
_F_T3 = "微软雅黑"
_F_T4 = "微软雅黑"
_F_LATIN = "Arial"
_SZ = {0: 26, 1: 16, 2: 14, 3: 12, 4: 10.5}

_COLOR_POS = RGBColor(0, 102, 102)      # 深青色：利多/正面
_COLOR_RISK = RGBColor(128, 0, 0)       # 暗红：利空/风险
_COLOR_NOTE = RGBColor(0x59, 0x59, 0x59)  # 灰：注释/来源
_PAGE_BG = "F2F2F2"                     # 浅灰页面
_HEADER_FILL = "1F3864"                 # 深蓝表头
_FORCE_LANDSCAPE = 6                    # 表格列数超过则横向

_NUM_RE = re.compile(
    r"(\d+(?:\.\d+)?\s*%|\d+(?:\.\d+)?\s*(?:万亿|亿元|万元|亿|元/股|元|万股)|\b\d{6}\b|\b(?:PE|PB)\s*\d+(?:\.\d+)?)"
)
_POS_KW = ["回购", "增持", "中标", "预增", "订单", "突破", "增长", "利好", "新高", "净流入", "上调", "超预期", "涨停", "签署", "翻倍"]
_RISK_KW = ["风险", "警示", "问询", "减持", "亏损", "立案", "违规", "利空", "下跌", "终止", "质押", "净流出", "下调", "被查", "爆雷", "提醒"]


def _style(run, sz, bold=False, color=None, underline=False, cn=_F_T3):
    run.font.name = _F_LATIN
    rPr = run._element.rPr
    rFonts = rPr.rFonts
    rFonts.set(qn("w:ascii"), _F_LATIN)
    rFonts.set(qn("w:hAnsi"), _F_LATIN)
    rFonts.set(qn("w:eastAsia"), cn)
    run.font.size = Pt(sz)
    run.font.bold = bold
    run.font.underline = underline
    if color is not None:
        run.font.color.rgb = color


def _tone(text):
    if any(k in text for k in _RISK_KW):
        return "risk"
    if any(k in text for k in _POS_KW):
        return "pos"
    return None


def _add_inline(par, text, sz, cn, tone):
    """正文：支持 **加粗**、核心数字(加粗+大一号)、利多/利空染色。"""
    for i, seg in enumerate(text.split("**")):
        if seg == "":
            continue
        base_bold = (i % 2 == 1)
        for part in _NUM_RE.split(seg):
            if part == "":
                continue
            if _NUM_RE.fullmatch(part):
                run = par.add_run(part)
                _style(run, sz + 1, bold=True, cn=cn)
            else:
                run = par.add_run(part)
                color, ul = None, False
                if tone == "pos":
                    color = _COLOR_POS
                elif tone == "risk":
                    color = _COLOR_RISK
                    ul = True
                _style(run, sz, bold=base_bold, cn=cn, color=color, underline=ul)


def _spacing(par, before, after, line=1.25):
    pf = par.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def _add_heading(doc, text, level):
    p = doc.add_paragraph()
    cn = _F_T0 if level == 0 else (_F_T1 if level == 1 else _F_T2)
    run = p.add_run(text)
    _style(run, _SZ[level], bold=True, cn=cn)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 0 else WD_ALIGN_PARAGRAPH.LEFT
    _spacing(p, 12 if level else 6, 6, line=1.0)
    return p


def _shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _cell_border(cell, edge="bottom", sz=8, color="000000"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:color"), color)
    borders.append(e)


def _style_table(table, header_fill=_HEADER_FILL):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, sz in (("top", 12), ("bottom", 12), ("left", 0), ("right", 0), ("insideH", 0), ("insideV", 0)):
        e = OxmlElement(f"w:{edge}")
        if sz == 0:
            e.set(qn("w:val"), "none")
        else:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)

    # 数值列判定
    ncols = len(table.columns)
    numeric_cols = set()
    for j in range(ncols):
        vals = []
        for row in table.rows[1:]:
            try:
                t = row.cells[j].text.strip()
            except Exception:
                continue
            if t in ("", "N/A", "-"):
                continue
            vals.append(t)
        if vals and all(re.fullmatch(r"[\d,.\-%+]+", v) for v in vals):
            numeric_cols.add(j)

    if table.rows:
        for cell in table.rows[0].cells:
            _shade_cell(cell, header_fill)
            _cell_border(cell, "bottom", 8)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _spacing(p, 2, 2, line=1.0)
                for run in p.runs:
                    _style(run, _SZ[4], bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for ri, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                _spacing(p, 2, 2, line=1.0)
                if j in numeric_cols and ri > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in p.runs:
                    _style(run, _SZ[4])
    if table.rows:
        for cell in table.rows[-1].cells:
            _cell_border(cell, "bottom", 12)


def _set_page(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    bg = OxmlElement("w:background")
    bg.set(qn("w:color"), _PAGE_BG)
    doc.element.insert(0, bg)
    settings = doc.settings.element
    if settings.find(qn("w:displayBackgroundShape")) is None:
        settings.append(OxmlElement("w:displayBackgroundShape"))


def _orient(doc, landscape):
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width


def _max_table_cols(lines):
    mx, i = 0, 0
    while i < len(lines):
        if lines[i].strip().startswith("|"):
            c = 0
            while i < len(lines) and lines[i].strip().startswith("|"):
                c = max(c, len(lines[i].strip().strip("|").split("|")))
                i += 1
            mx = max(mx, c)
        else:
            i += 1
    return mx


def build_docx(md_text, path, title):
    md_text = md_text or ""
    lines = md_text.splitlines()
    _orient(Document(), False)  # placeholder to keep import warm (no op)

    first_title = next((ln[2:].strip() for ln in lines if ln.startswith("# ")), None)
    doc_title = first_title or title or "报告"
    doc = Document()
    _set_page(doc)
    _orient(doc, _max_table_cols(lines) > _FORCE_LANDSCAPE)
    _add_heading(doc, doc_title, 0)

    i, skip_first_hash = 0, (first_title is not None)
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("# "):
            if skip_first_hash:
                skip_first_hash = False
            else:
                _add_heading(doc, ln[2:].strip(), 1)
        elif ln.startswith("## "):
            _add_heading(doc, ln[3:].strip(), 1)
        elif ln.startswith("### "):
            _add_heading(doc, ln[4:].strip(), 2)
        elif ln.startswith("|") and "|" in ln:
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [c.strip() for c in raw.split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                tbl = doc.add_table(rows=0, cols=max(len(r) for r in rows))
                for r in rows:
                    tcell = tbl.add_row().cells
                    for j, c in enumerate(r[: len(tcell)]):
                        tcell[j].text = c
                _style_table(tbl)
            continue
        elif ln.startswith(">"):
            p = doc.add_paragraph()
            _add_inline(p, ln.lstrip("> ").strip(), _SZ[4], _F_T4, None)
            for run in p.runs:
                run.font.color.rgb = _COLOR_NOTE
            _spacing(p, 2, 4, line=1.25)
        elif ln.startswith("- ") or ln.startswith("* "):
            text = ln[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, text, _SZ[3], _F_T3, _tone(text))
            _spacing(p, 0, 6, line=1.25)
        else:
            p = doc.add_paragraph()
            _add_inline(p, ln.strip(), _SZ[3], _F_T3, _tone(ln))
            _spacing(p, 0, 6, line=1.25)
        i += 1
    doc.save(str(path))


def md_to_docx(md_text, path, title):
    build_docx(md_text, path, title)